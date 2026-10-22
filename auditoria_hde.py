#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Módulo de Auditoría y Verificación de Datos - Hablemos de Economía (HDE)
Permite exportar e importar bidireccionalmente entre datos.json, control_publicaciones.csv
y auditoria_datos_hde.xlsx para auditorías asistidas por IA (Claude / Gemini) o revisión humana.

Incorpora el Escudo de Privacidad Institucional (Modo Low Profile):
- Sanitización estricta de identidad institucional.
- Columna automática CONTROL_ANONIMATO en Excel y CSV.
- Bloqueo absoluto de verificación si se detecta cualquier rastro personal.
"""

import sys
import os
import json
import csv
import argparse
import base64
import re

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

JSON_PATH = "datos.json"
CSV_PATH = "control_publicaciones.csv"
EXCEL_PATH = "auditoria_datos_hde.xlsx"

# -------------------------------------------------------------------------
# ESCUDO DE PRIVACIDAD: Patrones y Funciones de Anonimato Institucional
# -------------------------------------------------------------------------
_B64_PROHIBIDO = b"XGIoZ2FzcGFyfHRyZXBbYcOhXXxnYXNwaSg/OnRyZTk5fHRyZSk/KVxifGdhc3BhclwnP3M/XHMqbm90aW9u"
_B64_NOTION = b"Z2FzcGFyXCc/cz9ccypub3Rpb24="
_B64_AUTOR_POR = b"XGJwb3JccytnYXNwYXIoXHMrdHJlcFthw6FdKT9cYg=="
_B64_AUTOR_DE = b"XGJkZVxzK2dhc3Bhcihccyt0cmVwW2HDoV0pP1xi"
_B64_AUTOR_INV = b"XGJpbnZlc3RpZ2FjaVtvw7NdblxzK2RlXHMrZ2FzcGFyKFxzK3RyZXBbYcOhXSk/XGI="
_B64_NOM_APE = b"XGJnYXNwYXJccyt0cmVwW2HDoV1cYg=="
_B64_NOM = b"XGJnYXNwYXJcYg=="
_B64_USR = b"XGJnYXNwaXRyZSg5OSk/XGI="
_B64_NICK = b"XGJnYXNwaVxi"
_B64_TR1 = b"XGJ0cmVwYVxzK2FccytjYXNpXGI="
_B64_TR2 = b"XGJ0cmVwYVxzK2FsXGI="
_B64_TR3 = b"XGJ0cmVwYVxzK2FcYg=="
_B64_TR4 = b"XGJhbFxzK3RyZXBhclxzK2FsXGI="
_B64_TR5 = b"XGJhbFxzK3RyZXBhclxzK2FcYg=="
_B64_TR6 = b"XGJ0cmVwYXJcYg=="
_B64_TR7 = b"XGJ0cmVwW2HDoV1cYg=="

PATRON_IDENTIDAD_PROHIBIDA = re.compile(base64.b64decode(_B64_PROHIBIDO).decode("utf-8"), re.IGNORECASE)
PATRON_AUTOR_NOTION = re.compile(base64.b64decode(_B64_NOTION).decode("utf-8"), re.IGNORECASE)
PATRON_PRIMERA_PERSONA_SINGULAR = re.compile(
    r"\b(yo|conmigo|investigué|escribí|hice)\b|\b(mi|mis)\s+(canal|video|guion|investigaci[oó]n|redacci[oó]n|opini[oó]n|criterio)\b", 
    re.IGNORECASE
)

def sanitizar_identidad(texto: str) -> str:
    if not texto or not isinstance(texto, str):
        return texto

    texto = PATRON_AUTOR_NOTION.sub("Archivo Editorial HDE", texto)
    texto = re.sub(base64.b64decode(_B64_AUTOR_POR).decode("utf-8"), "por el Equipo Editorial HDE", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_AUTOR_DE).decode("utf-8"), "del Equipo Editorial HDE", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_AUTOR_INV).decode("utf-8"), "investigación del Equipo Editorial HDE", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_NOM_APE).decode("utf-8"), "Equipo Editorial HDE", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_NOM).decode("utf-8"), "Equipo Editorial HDE", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_USR).decode("utf-8"), "hablemosdeeco", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_NICK).decode("utf-8"), "hablemosdeeco", texto, flags=re.IGNORECASE)

    texto = re.sub(base64.b64decode(_B64_TR1).decode("utf-8"), "asciende a casi", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_TR2).decode("utf-8"), "asciende al", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_TR3).decode("utf-8"), "asciende a", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_TR4).decode("utf-8"), "al saltar al", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_TR5).decode("utf-8"), "al saltar a", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_TR6).decode("utf-8"), "escalar", texto, flags=re.IGNORECASE)
    texto = re.sub(base64.b64decode(_B64_TR7).decode("utf-8"), "asciende", texto, flags=re.IGNORECASE)

    texto = re.sub(r"\bmi\s+canal\b", "nuestro canal", texto, flags=re.IGNORECASE)
    texto = re.sub(r"\bmis\s+guiones\b", "nuestros informes", texto, flags=re.IGNORECASE)
    texto = re.sub(r"\bmi\s+investigaci[oó]n\b", "nuestra investigación", texto, flags=re.IGNORECASE)
    texto = re.sub(r"\binvestigu[eé]\b", "analizamos", texto, flags=re.IGNORECASE)
    texto = re.sub(r"\bmi\s+cuenta\b", "nuestra cuenta", texto, flags=re.IGNORECASE)

    return texto

def verificar_anonimato(texto_o_dict) -> tuple[bool, str]:
    if isinstance(texto_o_dict, dict):
        texto = " ".join(str(v) for v in texto_o_dict.values() if isinstance(v, str))
    else:
        texto = str(texto_o_dict or "")

    coincidencias_id = PATRON_IDENTIDAD_PROHIBIDA.findall(texto)
    if coincidencias_id:
        return False, f"Detectado nombre o término prohibido: {coincidencias_id}"

    coincidencias_sing = PATRON_PRIMERA_PERSONA_SINGULAR.findall(texto)
    if coincidencias_sing:
        return False, f"Detectada referencia en 1ra persona singular: {coincidencias_sing}"

    return True, "LIMPIO"

# -------------------------------------------------------------------------
# Estructura del Libro Excel de Auditoría
# -------------------------------------------------------------------------
COLUMNAS_EXCEL = [
    ("ID", 8, "center"),
    ("PUBLICADO", 12, "center"),
    ("VERIFICADO", 13, "center"),
    ("CONTROL_ANONIMATO", 22, "center"),
    ("ORIGEN_GUION", 24, "left"),
    ("FECHA_HISTORICA", 24, "left"),
    ("ES_EFEMERIDE", 14, "center"),
    ("CATEGORIA", 22, "left"),
    ("TITULAR_SLIDE_1", 45, "left"),
    ("LARGO_TITULAR", 16, "center"),
    ("DATO_CLAVE_SLIDE_2", 55, "left"),
    ("LARGO_DATO", 14, "center"),
    ("FUENTE_OFICIAL", 34, "left"),
    ("TIPO_VALOR_Y_MONEDA", 36, "left"),
    ("CAPTION_INSTAGRAM", 60, "left"),
    ("ESTADO_AUDITORIA_IA", 22, "center"),
    ("PROBLEMA_DETECTADO", 35, "left"),
    ("ARREGLO_SUGERIDO", 35, "left"),
]

def actualizar_csv_audit(posts, ruta_csv=CSV_PATH):
    """Sincroniza la planilla CSV de control con codificación utf-8-sig y delimitador ';'."""
    with open(ruta_csv, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow([
            "ID", "ESTADO", "VERIFICADO", "CONTROL_ANONIMATO", "ORIGEN_GUION", "FECHA HISTÓRICA",
            "TITULAR", "DATO CLAVE", "FUENTE OFICIAL", "TIPO DE VALOR",
            "TEXTO FINAL INSTAGRAM", "TEXTO FINAL X (TWITTER)"
        ])
        for p in posts:
            es_limpio, _ = verificar_anonimato(p)
            val_anon = "LIMPIO" if es_limpio else "⚠️ ALERTA IDENTIDAD"
            fecha = sanitizar_identidad(p.get("fecha_historica", ""))
            titular = sanitizar_identidad(p.get("titular", ""))
            dato = sanitizar_identidad(p.get("dato_clave", ""))
            caption_ig = sanitizar_identidad(p.get("caption", ""))
            fuente = sanitizar_identidad(p.get("fuente_oficial", ""))
            tipo_val = sanitizar_identidad(p.get("tipo_valor", ""))
            origen = sanitizar_identidad(p.get("origen_guion", "Lote Fundacional HDE"))
            texto_x = f"🏛️ {fecha.upper()}\n\n📌 {titular}\n\n{dato}\n\n#HistoriaEconomica #Argentina"
            estado = "PUBLICADO" if p.get("publicado") else "PENDIENTE"
            verificado = "SI" if (p.get("verificado") and es_limpio) else "NO"
            writer.writerow([
                p.get("id"), estado, verificado, val_anon, origen, fecha,
                titular, dato, fuente, tipo_val, caption_ig, texto_x
            ])
    print(f"Planilla CSV de auditoría ({ruta_csv}) sincronizada.")

def exportar_excel():
    """Genera auditoria_datos_hde.xlsx y sincroniza control_publicaciones.csv desde datos.json."""
    if not os.path.exists(JSON_PATH):
        print(f"Error: No se encontró {JSON_PATH}")
        sys.exit(1)

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        posts = json.load(f)

    # Sanitización previa en memoria de todos los campos de texto
    for p in posts:
        for k, v in list(p.items()):
            if isinstance(v, str):
                p[k] = sanitizar_identidad(v)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Auditoría HDE"
    ws.views.sheetView[0].showGridLines = True

    # Estilos de diseño institucional HDE
    header_fill = PatternFill(start_color="0C0E12", end_color="0C0E12", fill_type="solid")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    font_base = Font(name="Segoe UI", size=10)
    font_code = Font(name="Consolas", size=9)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_alerta = Font(name="Segoe UI", size=10, bold=True, color="991B1B")

    border_thin = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    fill_pub = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")     # Verde suave
    fill_verif = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")   # Amarillo suave
    fill_noverif = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # Rojo suave
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    # 1. Escribir Encabezados
    for col_idx, (col_nombre, col_width, _) in enumerate(COLUMNAS_EXCEL, start=1):
        cell = ws.cell(row=1, column=col_idx, value=col_nombre)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        ws.column_dimensions[get_column_letter(col_idx)].width = col_width

    ws.row_dimensions[1].height = 28

    total_limpios = 0
    total_alertas = 0

    # 2. Poblar datos
    for row_idx, p in enumerate(posts, start=2):
        es_pub = p.get("publicado", False)
        es_limpio, detalle_anon = verificar_anonimato(p)

        if es_limpio:
            val_anon = "LIMPIO"
            total_limpios += 1
        else:
            val_anon = "⚠️ ALERTA IDENTIDAD"
            total_alertas += 1

        # Si hay alerta de identidad, el post queda inhabilitado de verificación
        es_verif = p.get("verificado", False) and es_limpio

        val_pub = "SI" if es_pub else "NO"
        val_verif = "SI" if es_verif else "NO"
        origen = sanitizar_identidad(p.get("origen_guion", "Lote Fundacional HDE"))
        fecha = sanitizar_identidad(p.get("fecha_historica", ""))
        es_efe = "SI" if p.get("es_efemeride", False) else "NO"
        cat = sanitizar_identidad(p.get("categoria", "HISTORIA MONETARIA"))
        titular = sanitizar_identidad(p.get("titular", ""))
        dato = sanitizar_identidad(p.get("dato_clave", ""))
        fuente = sanitizar_identidad(p.get("fuente_oficial", ""))
        tipo_val = sanitizar_identidad(p.get("tipo_valor", ""))
        caption = sanitizar_identidad(p.get("caption", ""))
        estado_ia = sanitizar_identidad(p.get("estado_auditoria_ia", ""))
        problema = sanitizar_identidad(p.get("problema_detectado", ""))
        arreglo = sanitizar_identidad(p.get("arreglo_sugerido", ""))

        valores_fila = [
            p.get("id"),
            val_pub,
            val_verif,
            val_anon,
            origen,
            fecha,
            es_efe,
            cat,
            titular,
            len(titular),
            dato,
            len(dato),
            fuente,
            tipo_val,
            caption,
            estado_ia,
            problema,
            arreglo
        ]

        ws.row_dimensions[row_idx].height = 22

        for col_idx, val in enumerate(valores_fila, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.border = border_thin
            _, _, align_h = COLUMNAS_EXCEL[col_idx - 1]

            if col_idx == 1:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx == 2:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if es_pub:
                    cell.fill = fill_pub
            elif col_idx == 3:
                cell.font = font_bold
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.fill = fill_verif if es_verif else fill_noverif
            elif col_idx == 4:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if es_limpio:
                    cell.font = font_bold
                    cell.fill = fill_pub
                else:
                    cell.font = font_alerta
                    cell.fill = fill_noverif
            elif col_idx in (10, 12):
                cell.font = font_code
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif col_idx in (9, 11, 15):
                cell.font = font_base
                cell.alignment = Alignment(horizontal=align_h, vertical="center", wrap_text=False)
            else:
                cell.font = font_base
                cell.alignment = Alignment(horizontal=align_h, vertical="center")

            if col_idx not in (2, 3, 4) and row_idx % 2 == 0:
                cell.fill = fill_zebra

    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(EXCEL_PATH)
    print(f"✅ Excel de auditoría generado exitosamente: {EXCEL_PATH} ({len(posts)} registros)")
    print(f"   Estado de Anonimato: {total_limpios}/{len(posts)} LIMPIO ({total_alertas} alertas)")

    actualizar_csv_audit(posts)

def importar_excel():
    if not os.path.exists(EXCEL_PATH):
        print(f"Error: No se encontró {EXCEL_PATH}")
        sys.exit(1)

    if not os.path.exists(JSON_PATH):
        print(f"Error: No se encontró {JSON_PATH}")
        sys.exit(1)

    with open(JSON_PATH, "r", encoding="utf-8") as f:
        posts = json.load(f)

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws = wb.active

    posts_map = {p["id"]: p for p in posts}
    actualizados = 0
    nuevos_verificados = 0
    bloqueados_seguridad = 0

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue

        try:
            post_id = int(row[0])
        except (ValueError, TypeError):
            continue

        if post_id not in posts_map:
            continue

        p = posts_map[post_id]

        val_verif = str(row[2] or "").strip().upper()
        val_anon = str(row[3] or "").strip().upper()
        origen_guion = sanitizar_identidad(str(row[4] or "").strip())
        fecha_hist = sanitizar_identidad(str(row[5] or "").strip())
        es_efe_str = str(row[6] or "").strip().upper()
        categoria = sanitizar_identidad(str(row[7] or "").strip())
        titular = sanitizar_identidad(str(row[8] or "").strip())
        dato_clave = sanitizar_identidad(str(row[10] or "").strip())
        fuente_oficial = sanitizar_identidad(str(row[12] or "").strip())
        tipo_valor = sanitizar_identidad(str(row[13] or "").strip())
        caption = sanitizar_identidad(str(row[14] or "").strip())
        estado_ia = sanitizar_identidad(str(row[15] or "").strip().upper())
        problema = sanitizar_identidad(str(row[16] or "").strip())
        arreglo = sanitizar_identidad(str(row[17] or "").strip())

        if titular:
            p["titular"] = titular
        if dato_clave:
            p["dato_clave"] = dato_clave
        if fecha_hist:
            p["fecha_historica"] = fecha_hist
        if origen_guion:
            p["origen_guion"] = origen_guion
        if categoria:
            p["categoria"] = categoria
        if fuente_oficial:
            p["fuente_oficial"] = fuente_oficial
        if tipo_valor:
            p["tipo_valor"] = tipo_valor
        if caption:
            p["caption"] = caption

        p["es_efemeride"] = (es_efe_str == "SI")
        p["estado_auditoria_ia"] = estado_ia
        p["problema_detectado"] = problema
        p["arreglo_sugerido"] = arreglo

        es_anonimo_limpio, detalle = verificar_anonimato(p)
        if not es_anonimo_limpio or val_anon != "LIMPIO":
            p["verificado"] = False
            bloqueados_seguridad += 1
            print(f"⚠️ BLOQUEO DE PRIVACIDAD: Post #{post_id} no puede verificarse por rastro de identidad ({detalle}).")
        else:
            verif_previo = p.get("verificado", False)
            if val_verif == "SI" or estado_ia in ("OK", "APROBADO", "CORREGIDO"):
                p["verificado"] = True
                if not verif_previo:
                    nuevos_verificados += 1
            elif val_verif == "NO":
                p["verificado"] = False

        actualizados += 1

    with open(JSON_PATH, "w", encoding="utf-8") as f:
        json.dump(posts, f, ensure_ascii=False, indent=2)

    print(f"✅ Importación completada: {actualizados} registros sincronizados en {JSON_PATH}.")
    print(f"   Nuevos posts marcados como verificados: {nuevos_verificados}")
    if bloqueados_seguridad > 0:
        print(f"   ⚠️ Posts bloqueados de verificación por alerta de identidad: {bloqueados_seguridad}")

    actualizar_csv_audit(posts)

def main():
    parser = argparse.ArgumentParser(description="Auditoría y Verificación de Datos HDE con Escudo de Privacidad")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--exportar", action="store_true", help="Exportar datos.json a auditoria_datos_hde.xlsx y CSV")
    group.add_argument("--importar", action="store_true", help="Importar correcciones de auditoria_datos_hde.xlsx a datos.json y CSV")

    args = parser.parse_args()

    if args.exportar:
        exportar_excel()
    elif args.importar:
        importar_excel()

if __name__ == "__main__":
    main()
